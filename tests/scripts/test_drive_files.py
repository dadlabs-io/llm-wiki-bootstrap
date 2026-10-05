#!/usr/bin/env python3
"""Checks that the Drive step reads every kind of file it is given (task #69, 2026-10-05): a PDF, an Office
file, a Google Slides deck and an image are saved as raws and queued with them; a text file with no link and a
type the cycle does not read are named in the report and left in Drive; nothing is skipped silently. Never
shipped; no Drive access: a fake service serves the files, and the real PDF fetcher, MarkItDown and queue
script run against a throwaway notebook.

    uv run python tests/scripts/test_drive_files.py    # exit 0 = every check passed
"""
from __future__ import annotations

import contextlib
import importlib.util
import io
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

results: list[tuple[bool, str]] = []


def check(name: str, ok: bool, detail: object = ""):
    results.append((bool(ok), name + (f"  [{detail}]" if not ok and detail != "" else "")))


spec = importlib.util.spec_from_file_location("wiki_fetch_drive_folder", ROOT / "scripts" / "wiki-fetch-drive-folder.py")
drive = importlib.util.module_from_spec(spec)
spec.loader.exec_module(drive)


def tiny_pdf(text: str) -> bytes:
    """A one-page PDF whose text pdftotext / pypdf can read."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        body = (f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream") if o is None else o.encode()
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def tiny_pptx(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Deck title"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


PNG = bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                    "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")
FILES = {
    "f-link": ("link.txt", "text/plain", b"A good article https://example.org/article-one"),
    "f-note": ("note.txt", "text/plain", b"remember to look at agent memory later"),
    "f-pptx": ("Agent Patterns.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation",
               tiny_pptx("Orchestrator worker pattern explained")),
    "f-gslides": ("Planning deck", "application/vnd.google-apps.presentation", None),
    "f-pdf": ("Agentic_Design_Patterns.pdf", "application/pdf", tiny_pdf("Agentic design patterns PDF text")),
    "f-png": ("Screenshot 2026-10-02.png", "image/png", PNG),
    "f-mp3": ("voice memo.mp3", "audio/mpeg", b"ID3"),
}
EXPORTED = tiny_pptx("Exported Google Slides text")


class _Call:
    def __init__(self, v):
        self.v = v

    def execute(self):
        return self.v


class FakeDrive:
    def files(self):
        return self

    def list(self, q="", **_kw):
        return _Call({"files": [{"id": k, "name": n, "mimeType": m, "modifiedTime": "2026-10-01T00:00:00Z"}
                                for k, (n, m, _b) in FILES.items()]})

    def get_media(self, fileId):  # noqa: N803 - the Google client's argument name
        return _Call(FILES[fileId][2])

    def export(self, fileId, mimeType):  # noqa: N803
        return _Call(EXPORTED if FILES[fileId][1] == "application/vnd.google-apps.presentation" else b"")


# Pure rules.
check("a text file or Google Doc is a link capture", drive.file_kind("a.txt", "text/plain") == "link"
      and drive.file_kind("Doc", "application/vnd.google-apps.document") == "link")
check("a PDF, an image, an Office file and Google Slides are files to save",
      [drive.file_kind(n, m) for n, m in (("x.pdf", "application/pdf"), ("s.png", "image/png"),
                                          ("d.docx", "application/octet-stream"),
                                          ("g", "application/vnd.google-apps.presentation"))] == ["pdf", "image", "document", "document"])
check("an mp3 or a zip is not read", drive.file_kind("v.mp3", "audio/mpeg") is None and drive.file_kind("a.zip", "application/zip") is None)

with tempfile.TemporaryDirectory() as td:
    tmp = Path(td)
    nb = tmp / "nb"
    (nb / "wiki").mkdir(parents=True)
    (nb / "_inbox" / "pending").mkdir(parents=True)
    moved_ids = []
    drive.get_drive_service = lambda *a, **k: FakeDrive()
    drive.find_folder_id = lambda service, name, parent_id=None: f"id-{name}"
    drive.resolve_url = lambda url, timeout=25: url
    drive.move_handled_files = lambda service, qr, entries, *a, **k: (
        moved_ids.extend(fid for fid, st, _m in qr if st in ("queued", "known")) or (len(moved_ids), 0, []))
    drive._write_activity_log = lambda entry: None
    out = tmp / "run" / "drive-fetch.md"
    argv, sys.argv = sys.argv, ["wiki-fetch-drive-folder.py", "--subfolder", "nb", "--queue-into", "nb",
                                "--queue-vault", str(tmp), "--archive-subfolder", "2026-10-05-01", "--out", str(out)]
    so, se = io.StringIO(), io.StringIO()
    try:
        with contextlib.redirect_stdout(so), contextlib.redirect_stderr(se):
            rc = drive.main()
    finally:
        sys.argv = argv
    check("the scan exits 0", rc == 0, se.getvalue()[-500:])
    raws = {p.name: p for p in (nb / "raw").glob("*")} if (nb / "raw").is_dir() else {}
    tickets = [p.read_text(encoding="utf-8") for p in (nb / "_inbox" / "pending").glob("*.md") if not p.name.startswith("_")]
    by_source = {next((ln.split(":", 1)[1].strip() for ln in t.splitlines() if ln.startswith("source:")), ""): t for t in tickets}

    def raw_of(fid):
        t = by_source.get(drive.drive_link(fid), "")
        line = next((ln for ln in t.splitlines() if ln.startswith("raw_path:")), "")
        p = nb / line.split(":", 1)[1].strip() if line else None
        return p.read_text(encoding="utf-8") if p and p.is_file() else ""

    check("the link file is queued by its URL, as before", "https://example.org/article-one" in by_source, list(by_source))
    pdf = raw_of("f-pdf")
    check("the PDF is saved as a raw through the PDF fetcher and queued with it",
          "Agentic design patterns PDF text" in pdf and any(n.endswith(".pdf") for n in raws), (pdf[:200], list(raws)))
    deck = raw_of("f-pptx")
    check("the PowerPoint is converted (MarkItDown) and queued with its raw",
          "Orchestrator worker pattern explained" in deck and "converted: markdown" in deck, deck[:300])
    check("... its original kept beside the raw", any(n.endswith(".pptx") for n in raws), list(raws))
    gs = raw_of("f-gslides")
    check("a Google Slides deck is exported and converted", "Exported Google Slides text" in gs, gs[:300])
    img = raw_of("f-png")
    check("the image is saved, and its raw tells the session to read it",
          "read it before writing the entry" in img and any(n.endswith(".png") for n in raws), img[:300])
    check("each file's ticket names its Drive link as source and carries raw_path",
          all(drive.drive_link(f) in by_source and "raw_path: raw/" in by_source[drive.drive_link(f)]
              for f in ("f-pdf", "f-pptx", "f-gslides", "f-png")), list(by_source))
    check("the note with no link and the mp3 are not queued",
          drive.drive_link("f-note") not in by_source and drive.drive_link("f-mp3") not in by_source)
    report = out.read_text(encoding="utf-8") if out.is_file() else ""
    check("the report names what was left in Drive and why",
          "## Left in Drive" in report and "`note.txt`: left in Drive: no link in it" in report
          and "`voice memo.mp3`: left in Drive: audio/mpeg is not a type the cycle reads" in report, report[-700:])
    check("the report lists the files and their result", "## Files (the source itself)" in report
          and "Agentic_Design_Patterns.pdf | pdf | queued" in report, report[-900:])
    step = json.loads(out.with_suffix(".json").read_text(encoding="utf-8")) if out.with_suffix(".json").is_file() else {}
    check("the step JSON queues the files by their Drive links",
          {drive.drive_link(f) for f in ("f-pdf", "f-pptx", "f-gslides", "f-png")} <= {q["url"] for q in step.get("queued", [])},
          step.get("queued"))
    check("... and defers the two left in Drive, with the reason",
          sorted(d["url"] for d in step.get("deferred", [])) == sorted([drive.drive_link("f-note"), drive.drive_link("f-mp3")])
          and step.get("summary", {}).get("files_left") == 2 and step.get("summary", {}).get("files_queued") == 4, step.get("deferred"))
    check("handled files move to _completed; the two left stay",
          {"f-pdf", "f-pptx", "f-gslides", "f-png", "f-link"} <= set(moved_ids)
          and "f-note" not in moved_ids and "f-mp3" not in moved_ids, moved_ids)

failed = [n for ok, n in results if not ok]
for ok, n in results:
    print(("PASS " if ok else "FAIL ") + n)
print(f"\n{len(results) - len(failed)}/{len(results)} passed")
sys.exit(1 if failed else 0)
